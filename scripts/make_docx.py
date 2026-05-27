"""Convert BAI_LAM_TIENG_VIET.md to a nicely formatted .docx file."""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "docs" / "research" / "BAI_LAM_TIENG_VIET.md"
DST = ROOT / "docs" / "research" / "BAI_LAM_TIENG_VIET.docx"


def set_cell_border(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single")
        b.set(qn("w:sz"), "4")
        b.set(qn("w:color"), "808080")
        tc_borders.append(b)
    tc_pr.append(tc_borders)


def shade_cell(cell, color="DDDDDD"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tc_pr.append(shd)


def add_runs_with_bold(paragraph, text):
    """Render simple inline markdown (**bold**, *italic*) inside a paragraph."""
    import re

    # Replace links [text](url) -> text (url)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)

    # Tokenize bold (**...**) and italic (*...*)
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        token = m.group()
        if token.startswith("**"):
            r = paragraph.add_run(token[2:-2])
            r.bold = True
        else:
            r = paragraph.add_run(token[1:-1])
            r.italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def build_doc(md_text: str) -> Document:
    doc = Document()

    # Default font: Times New Roman 12 (set on Normal style)
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    rpr = normal.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts")) or OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:cs"), "Times New Roman")
    rpr.append(rfonts)

    # Page margins ~ 2.5cm
    for section in doc.sections:
        section.top_margin = Cm(2.5)
        section.bottom_margin = Cm(2.5)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()

        # Skip horizontal rules and blank lines
        if line.strip() in ("", "---"):
            i += 1
            continue

        # Heading 1
        if line.startswith("# "):
            p = doc.add_heading(level=1)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_runs_with_bold(p, line[2:].strip())
            i += 1
            continue
        if line.startswith("## "):
            p = doc.add_heading(level=2)
            add_runs_with_bold(p, line[3:].strip())
            i += 1
            continue
        if line.startswith("### "):
            p = doc.add_heading(level=3)
            add_runs_with_bold(p, line[4:].strip())
            i += 1
            continue

        # Markdown table
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].lstrip().startswith("|---"):
            header_cells = [c.strip() for c in line.strip().strip("|").split("|")]
            i += 2  # skip separator
            rows = []
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                row_cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                rows.append(row_cells)
                i += 1
            table = doc.add_table(rows=1 + len(rows), cols=len(header_cells))
            table.autofit = True
            # header
            for c, txt in enumerate(header_cells):
                cell = table.rows[0].cells[c]
                cell.text = ""
                p = cell.paragraphs[0]
                add_runs_with_bold(p, txt)
                for run in p.runs:
                    run.bold = True
                shade_cell(cell, "E6E6E6")
                set_cell_border(cell)
            # body
            for r, row in enumerate(rows, start=1):
                for c, txt in enumerate(row):
                    cell = table.rows[r].cells[c]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    add_runs_with_bold(p, txt)
                    set_cell_border(cell)
            doc.add_paragraph()
            continue

        # Bulleted list
        if line.lstrip().startswith(("- ", "* ")):
            content = line.lstrip()[2:]
            p = doc.add_paragraph(style="List Bullet")
            add_runs_with_bold(p, content)
            i += 1
            continue

        # Numbered list (1. ...)
        import re
        m = re.match(r"^\s*(\d+)\.\s+(.*)$", line)
        if m:
            content = m.group(2)
            p = doc.add_paragraph(style="List Number")
            add_runs_with_bold(p, content)
            i += 1
            continue

        # Regular paragraph (may span multiple lines until blank)
        para_lines = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].lstrip().startswith(
            ("#", "|", "-", "*", "1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.")
        ):
            para_lines.append(lines[i].rstrip())
            i += 1
        para_text = " ".join(para_lines).strip()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_runs_with_bold(p, para_text)

    return doc


def main():
    md = SRC.read_text(encoding="utf-8")
    doc = build_doc(md)
    doc.save(DST)
    print(f"Wrote: {DST}  ({DST.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
